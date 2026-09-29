#!/usr/bin/env python3
"""Statically read table metadata; never execute the supplied JavaScript."""
from __future__ import annotations
import argparse
from copy import copy
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from urllib.parse import parse_qs, urlsplit

import esprima
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'assets' / 'import-fields-template.xlsx'
HEADERS = ['字段编号', '字段名称', '模块编号', '模块名称', '字段顺序', '字段加密类型',
           '字段加密长度', '字段加密开始位置', '查询字段名称', '所属字段编号', '所属字段名称',
           '背景色', '加密排除类型', '展示类型', '注意事项说明']

class InputError(ValueError):
    pass

class Page(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.scripts = []
        self.headers = []
        self.script = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'script' and not attrs.get('src'):
            kind = attrs.get('type', '')
            if kind in ('', 'text/javascript', 'application/javascript', 'module'):
                self.script = []
        for key in ('th:include', 'th:replace', 'data-th-include', 'data-th-replace'):
            value = attrs.get(key, '')
            if re.search(r'\binclude\s*::\s*header\s*\(', value):
                self.headers.append(value)

    def handle_endtag(self, tag):
        if tag == 'script' and self.script is not None:
            self.scripts.append(''.join(self.script))
            self.script = None

    def handle_data(self, data):
        if self.script is not None:
            self.script.append(data)


def walk(node):
    if isinstance(node, dict):
        yield node
        for value in node.values():
            yield from walk(value)
    elif isinstance(node, list):
        for value in node:
            yield from walk(value)


def literal(node, label):
    if node and node.get('type') == 'Literal' and isinstance(node.get('value'), str):
        return node['value']
    if node and node.get('type') == 'TemplateLiteral' and not node['expressions']:
        return ''.join(q['value']['cooked'] for q in node['quasis'])
    raise InputError(f'{label} 必须是静态字符串，不能自动推测动态表达式。')


def key_of(prop):
    if prop.get('type') != 'Property' or prop.get('computed'):
        return None
    key = prop['key']
    return key.get('name') if key['type'] == 'Identifier' else key.get('value')


def module_name(page):
    names = set()
    for header in page.headers:
        match = re.search(r'\bheader\s*\(\s*((?:\'(?:\\.|[^\'\\])*\')|(?:"(?:\\.|[^"\\])*"))\s*\)', header)
        if not match:
            raise InputError('header(...) 的模块名称不是可识别的静态字符串。')
        names.add(literal(esprima.parseScript(match[1]).toDict()['body'][0]['expression'], '模块名称'))
    if len(names) != 1 or not next(iter(names), '').strip():
        raise InputError('需要唯一且非空的 include :: header(模块名称)。')
    return names.pop()


def extract(html):
    page = Page()
    page.feed(html)
    name = module_name(page)
    trees = []
    for script in page.scripts:
        try:
            trees.append(esprima.parseScript(script).toDict())
        except Exception as exc:
            if re.search(r'\bcolumns\b|tplCd', script):
                raise InputError(f'包含 columns/tplCd 的脚本无法静态解析：{exc}') from exc
    arrays, codes = [], set()
    for tree in trees:
        for node in walk(tree):
            if node.get('type') == 'Property' and key_of(node) == 'columns':
                if node['value']['type'] != 'ArrayExpression':
                    raise InputError('columns 必须直接定义为数组；变量引用或动态生成需要人工确认。')
                arrays.append(node['value'])
            if node.get('type') in ('Literal', 'TemplateLiteral'):
                try:
                    value = literal(node, '导出地址')
                except InputError:
                    continue
                if re.search(r'/export\?', value):
                    params = parse_qs(urlsplit(value).query, keep_blank_values=True)
                    for code in params.get('tplCd', []):
                        if not re.fullmatch(r'[A-Za-z0-9_.-]+', code):
                            raise InputError('export?tplCd= 的模块编号为空或不受支持。')
                        codes.add(code)
    if len(codes) != 1:
        raise InputError('需要唯一的 /export?tplCd=模块编号；缺失或多个编号时不猜测。')
    if len(arrays) != 1:
        raise InputError(f'需要唯一的 columns 数组，实际找到 {len(arrays)} 个。')
    columns = []
    def collect(array):
        for node in array['elements']:
            if not node:
                raise InputError('columns 中包含空项。')
            if node['type'] == 'ArrayExpression':
                collect(node)
                continue
            if node['type'] != 'ObjectExpression':
                raise InputError('columns 中存在动态列或展开语法，无法保证完整顺序。')
            props = {}
            for prop in node['properties']:
                key = key_of(prop)
                if prop.get('type') == 'SpreadElement':
                    raise InputError('列对象存在展开语法，可能覆盖 field/title。')
                if key not in ('field', 'title'):
                    continue
                if key in props:
                    raise InputError(f'列属性重复：{key}')
                props[key] = prop['value']
            title = literal(props.get('title'), '列 title')
            if title.strip() == '序号':
                continue
            field = literal(props.get('field'), f'列“{title}”的 field')
            if not field.strip() or not title.strip():
                raise InputError('field/title 不能为空。')
            columns.append((field, title))
    collect(arrays[0])
    if not columns:
        raise InputError('排除序号后没有可导出的字段。')
    if len({field for field, _ in columns}) != len(columns):
        raise InputError('存在重复 field，需先确认 HTML 列定义。')
    return codes.pop(), name, columns


def filename(name):
    if re.search(r'[\x00-\x1f<>:"/\\|?*]', name) or name in ('.', '..') or name.endswith((' ', '.')):
        raise InputError('模块名称包含文件名不允许的字符，无法原样另存为。')
    if len((name + '.xlsx').encode('utf-8')) > 240:
        raise InputError('模块名称过长，无法安全生成文件名。')
    return name + '.xlsx'


def export(source, output_dir=None):
    source = Path(source).expanduser()
    if not source.is_absolute():
        raise InputError('请输入 HTML 的绝对路径。')
    if not source.is_file():
        raise InputError(f'HTML 不存在：{source}')
    code, name, columns = extract(source.read_text(encoding='utf-8-sig'))
    output_dir = Path(output_dir).expanduser() if output_dir else Path.home() / 'Desktop'
    if not output_dir.is_dir():
        raise InputError(f'输出目录不存在：{output_dir}')
    target = output_dir / filename(name)
    if target.exists() or target.is_symlink():
        raise InputError(f'目标已存在，未覆盖：{target}；请先移动或重命名现有文件。')
    book = openpyxl.load_workbook(TEMPLATE)
    sheet = book.active
    if [sheet.cell(1, i).value for i in range(1, 16)] != HEADERS:
        raise InputError('内置模板表头不匹配。')
    styles = [copy(sheet.cell(3, i)._style or book._cell_styles[0]) for i in range(1, 16)]
    height = sheet.row_dimensions[3].height
    # Keep rows 1–2 (including the original instructions) unchanged.
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.value = None
    for index, (field, title) in enumerate(columns, 1):
        values = [field, title, code, name, index, 0, 0, 0, field, None, None, None, 1, 0, None]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index + 2, col)
            cell._style = copy(styles[col - 1])
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'  # Never interpret input text as an Excel formula.
        sheet.row_dimensions[index + 2].height = height
    if sheet.max_row > len(columns) + 2:
        sheet.delete_rows(len(columns) + 3, sheet.max_row - len(columns) - 2)
    # Complete the workbook before publishing; hard-link refuses concurrent overwrite.
    fd, temporary = tempfile.mkstemp(prefix='.fields-', suffix='.xlsx', dir=output_dir)
    os.close(fd)
    try:
        book.save(temporary)
        check = openpyxl.load_workbook(temporary)
        if check.active.cell(len(columns) + 2, 1).value != columns[-1][0]:
            raise InputError('生成后回读校验失败。')
        check.close()
        os.link(temporary, target)
    finally:
        book.close()
        Path(temporary).unlink(missing_ok=True)
    return {'output': str(target.resolve()), 'module_code': code, 'module_name': name,
            'field_count': len(columns), 'fields': [field for field, _ in columns]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('html', help='HTML absolute path')
    parser.add_argument('--output-dir', help='Override Desktop (primarily for tests)')
    args = parser.parse_args()
    try:
        print(json.dumps(export(args.html, args.output_dir), ensure_ascii=False, indent=2))
    except (InputError, OSError, UnicodeError) as exc:
        print(f'错误：{exc}', file=sys.stderr)
        return 2
    return 0

if __name__ == '__main__':
    sys.exit(main())

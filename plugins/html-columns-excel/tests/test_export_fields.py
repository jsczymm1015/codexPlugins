from pathlib import Path
from copy import copy
import sys
import tempfile
import unittest
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from export_fields import extract, export, InputError, TEMPLATE


def page(columns, extra='', name='测试模块', code='exampleList'):
    return f'''<head><th:block th:include="include :: header('{name}')"/></head>
    <script>const url='/export?tplCd={code}'; const options={{columns: {columns}}};{extra}</script>'''

NORMAL = """[{title:'序号',formatter:function(v,r,i){return i+1;}},
{field:'id', title:'编号'}, {field:'reason',title:'原因',visible:false,
formatter:function(v){const obj={title:'伪列',field:'wrong'};return /[},]/.test(v)?obj.title:v;}}]"""

class Extraction(unittest.TestCase):
    def test_order_hidden_formatter_and_sequence(self):
        self.assertEqual(extract(page(NORMAL)), ('exampleList','测试模块',[('id','编号'),('reason','原因')]))
    def test_ignore_other_properties(self):
        self.assertEqual(extract(page("[{field:'a',title:'甲',visible:check(),formatter:function(v){return v===1?'是':'否';}}]"))[2],[('a','甲')])
    def test_comments_not_columns(self):
        self.assertEqual(len(extract(page(NORMAL, "// columns: [{field:'fake',title:'假'}]\n"))[2]),2)
    def test_nested_arrays(self):
        self.assertEqual(extract(page("[[{field:'a',title:'甲'}],[{field:'b',title:'乙'}]]"))[2], [('a','甲'),('b','乙')])
    def test_escaped_and_template_strings(self):
        self.assertEqual(extract(page(r'''[{field:`a`,title:'A\'B'}]'''))[2],[('a',"A'B")])
    def test_query_order(self):
        self.assertEqual(extract(page(NORMAL).replace('/export?tplCd=', '/export?other=1&tplCd='))[0], 'exampleList')
    def test_repeated_same_module(self):
        self.assertEqual(extract(page(NORMAL,"const second='/export?tplCd=exampleList';"))[0], 'exampleList')
    def test_ambiguities_and_unsupported_fail(self):
        cases = [page(NORMAL,"const other={columns:[]};"), page(NORMAL,"const other='/export?tplCd=other';"),
                 page('buildColumns()'), page("[{title:'操作'}]"), page("[{field:'a',title:getTitle()}]"),
                 page("[{field:'a',title:'甲'},{field:'a',title:'乙'}]"), page('[]'),
                 page("[{...other,field:'a',title:'甲'}]"), page(NORMAL,code=''),
                 page(NORMAL).replace('tplCd=', 'other='), page(NORMAL).replace('include :: header', 'include :: footer'),
                 page(NORMAL).replace("const options", "const broken = ; const options")]
        for html in cases:
            with self.subTest(html=html), self.assertRaises(InputError):
                extract(html)

class Workbook(unittest.TestCase):
    def test_values_styles_template_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            html = root/'list.html'
            html.write_text(page(NORMAL), encoding='utf-8')
            result = export(html,root)
            self.assertEqual(result['field_count'],2)
            path=root/'测试模块.xlsx'
            book=openpyxl.load_workbook(path)
            template=openpyxl.load_workbook(TEMPLATE)
            sheet=book.active
            self.assertEqual(sheet.max_row,4)
            for row in (1,2):
                self.assertEqual([c.value for c in sheet[row]], [c.value for c in template.active[row]])
            self.assertEqual([c.value for c in sheet[3]],['id','编号','exampleList','测试模块',1,0,0,0,'id',None,None,None,1,0,None])
            self.assertEqual(sheet['A4'].value,'reason')
            self.assertEqual(sheet['E4'].value,2)
            self.assertEqual(sheet['A3']._style,template.active['A3']._style or template._cell_styles[0])
            self.assertEqual(sheet.column_dimensions['B'].width,template.active.column_dimensions['B'].width)
            book.close();template.close()
            original=path.read_bytes()
            with self.assertRaises(InputError): export(html,root)
            self.assertEqual(path.read_bytes(),original)
    def test_formula_stored_as_text(self):
        with tempfile.TemporaryDirectory() as temp:
            html=Path(temp)/'list.html'
            html.write_text(page("[{field:'=1+1',title:'=2+2'}]"),encoding='utf-8')
            result=export(html,temp)
            book=openpyxl.load_workbook(result['output'])
            self.assertEqual(book.active['A3'].data_type,'s')
            self.assertEqual(book.active['B3'].value,'=2+2')
            book.close()
    def test_invalid_filename_no_output(self):
        with tempfile.TemporaryDirectory() as temp:
            html=Path(temp)/'list.html'
            html.write_text(page(NORMAL,name='../bad'),encoding='utf-8')
            with self.assertRaises(InputError): export(html,temp)
            self.assertEqual(list(Path(temp).glob('*.xlsx')),[])
    def test_absolute_path_required(self):
        with self.assertRaises(InputError): export('list.html')
    def test_more_than_template_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            html=Path(temp)/'list.html'
            html.write_text(page('['+','.join("{field:'f%d',title:'列%d'}"%(i,i) for i in range(30))+']'),encoding='utf-8')
            result=export(html,temp)
            book=openpyxl.load_workbook(result['output'])
            self.assertEqual(book.active.max_row,32)
            self.assertEqual(book.active['E32'].value,30)
            self.assertEqual(book.active['A32']._style,book.active['A3']._style)
            book.close()

if __name__=='__main__': unittest.main()
